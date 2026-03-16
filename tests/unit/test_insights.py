from pathlib import Path

from docx import Document
from fpdf import FPDF

from services.insights import extract_metadata


def test_extract_metadata_txt(tmp_path: Path) -> None:
    """
    Validates the extraction of metadata from .txt files.

    Args:
        tmp_path (Path): Temporary path provided by pytest.
    Returns:
        None.
    Raises:
        AssertionError: If the metadata isn't extracted correctly.
    """

    file = tmp_path / "test.txt"
    file.write_text("Dit is een Nederlands testdocument.")

    metadata = extract_metadata(file)

    assert metadata["filename"] == "test.txt"
    assert metadata["file_type"] == ".txt"
    assert metadata["word_count"] == 5
    assert metadata["file_size"] > 0
    assert metadata["language"] == "nl"


def test_extract_metadata_docx(tmp_path: Path) -> None:
    """
    Validates the extraction of metadata from .docx files.

    Args:
        tmp_path (Path): Temporary path provided by pytest.
    Returns:
        None.
    Raises:
        AssertionError: If the metadata isn't extracted correctly.
    """

    docx_file = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("This is a test document")
    doc.save(docx_file)

    metadata = extract_metadata(docx_file)

    assert metadata["filename"] == "test.docx"
    assert metadata["file_type"] == ".docx"
    assert metadata["word_count"] == 5
    assert metadata["file_size"] > 0
    assert metadata["language"] == "en"
    assert metadata["paragraph_count"] == 1
    assert metadata["table_count"] == 0


def test_extract_metadata_pdf(tmp_path: Path) -> None:
    """
    Validates the extraction of metadata from .pdf files.

    Args:
        tmp_path (Path): Temporary path provided by pytest.
    Returns:
        None.
    Raises:
        AssertionError: If the metadata isn't extracted correctly.
    """

    pdf_file = tmp_path / "test.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(0, 10, "This is a test document")
    pdf.output(str(pdf_file))

    metadata = extract_metadata(pdf_file)

    assert metadata["filename"] == "test.pdf"
    assert metadata["file_type"] == ".pdf"
    assert metadata["word_count"] == 5
    assert metadata["file_size"] > 0
    assert metadata["language"] == "en"
